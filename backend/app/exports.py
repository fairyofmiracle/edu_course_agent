from __future__ import annotations
from zipfile import ZIP_DEFLATED, ZipFile

import json
import re
from io import BytesIO
from typing import Any

from .legal_basis import required_legal_basis
from .rationale import module_rationales
from .models import CourseDocument

FORMAT_LABELS = {
    "lecture": "очная",
    "blended": "очно-заочная",
    "workshop": "практическая",
}
KIND_MEDIA = {
    "json": "application/json",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "zip": "application/zip",
    "moodle": "application/xml",
    "gift": "text/plain; charset=utf-8",
    "csv": "text/csv; charset=utf-8",
    "html": "text/html; charset=utf-8",
}


def hours(value: float) -> str:
    return str(round(value * 2) / 2).replace(".", ",")


def artifact_filename(course: CourseDocument, kind: str) -> str:
    base = re.sub(r"[^\w\-]+", "-", (course.brief.title or "course").strip(), flags=re.UNICODE)
    base = base.strip("-")[:60] or "course"
    if kind == "moodle":
        return f"{base}-moodle.xml"
    if kind == "gift":
        return f"{base}.gift.txt"
    if kind == "csv":
        return f"{base}-tests.csv"
    if kind == "html":
        return f"{base}-slides.html"
    return f"{base}.{kind}"


#BytesIO() создаёт буфер в оперативной памяти
#ZipFile() формирует архив
# writestr() добавляет в него файлы.
def _file_stem(value: str, fallback: str) -> str:
    base = re.sub(r"[^\w\-]+", "-", (value or "").strip(), flags=re.UNICODE)
    return base.strip("-")[:50] or fallback


def build_zip(course: CourseDocument) -> bytes:
    """Общая выгрузка: программа, лекции, тесты и слайды из одного JSON."""
    from .lms_export import course_csv, course_gift, course_moodle_xml, course_slides_html

    buffer = BytesIO()
    title = _file_stem(course.brief.title, "programma")
    module_titles = {module.id: module.title for module in course.modules}

    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(f"основной документ/{title}.docx", build_docx(course))
        for index, lecture in enumerate(course.lectures, start=1):
            module_title = module_titles.get(lecture.module_id, lecture.module_id)
            stem = f"{index:02d}-{_file_stem(module_title, lecture.module_id)}"
            archive.writestr(f"презентации/{stem}.pptx", build_lecture_pptx(course, lecture))
            archive.writestr(f"лекции/{stem}.docx", build_lecture_docx(course, lecture))
        archive.writestr(f"тесты/{title}.gift.txt", course_gift(course).encode("utf-8"))
        archive.writestr(f"тесты/{title}-moodle.xml", course_moodle_xml(course).encode("utf-8"))
        archive.writestr(f"тесты/{title}.csv", course_csv(course).encode("utf-8"))
        archive.writestr(f"слайды/{title}.html", course_slides_html(course).encode("utf-8"))
        from .audit import audit_course_document

        report = audit_course_document(course)
        archive.writestr(
            f"проверка/{title}-audit.json",
            json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8"),
        )

    return buffer.getvalue()


def build_artifact(kind: str, course: CourseDocument) -> tuple[bytes, str, str]:
    fmt = kind.lower().strip()

    if fmt not in KIND_MEDIA:
        raise ValueError("Доступны форматы json, docx, pptx, zip, moodle, gift, csv и html")

    from .lms_export import course_csv, course_gift, course_moodle_xml, course_slides_html

    if fmt == "json":
        payload = json.dumps(course.as_api(), ensure_ascii=False, indent=2).encode("utf-8")
    elif fmt == "docx":
        payload = build_docx(course)
    elif fmt == "pptx":
        payload = build_pptx(course)
    elif fmt == "zip":
        payload = build_zip(course)
    elif fmt == "moodle":
        payload = course_moodle_xml(course).encode("utf-8")
    elif fmt == "gift":
        payload = course_gift(course).encode("utf-8")
    elif fmt == "csv":
        payload = course_csv(course).encode("utf-8")
    else:
        payload = course_slides_html(course).encode("utf-8")

    return payload, KIND_MEDIA[fmt], artifact_filename(course, fmt)


def _program_title(kind: str) -> str:
    if kind == "qualification":
        return "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА ПОВЫШЕНИЯ КВАЛИФИКАЦИИ"
    if kind == "general":
        return "ДОПОЛНИТЕЛЬНАЯ ОБЩЕОБРАЗОВАТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА"
    return "ДОПОЛНИТЕЛЬНАЯ ОБРАЗОВАТЕЛЬНАЯ ПРОГРАММА"


def _source_mark(source: str | None, title: str | None = None) -> str:
    """Пометка источника для выгрузки: официальная / файл / не подтверждён."""
    from .legal_basis import is_official_publication
    from .verbatim import stored_source_status

    raw = (source or "").strip()
    if is_official_publication(raw):
        return "официальная публикация"
    status = stored_source_status(raw)
    if status == "source_attached":
        return "файл методиста"
    lowered = raw.lower()
    if lowered.startswith("http://") or lowered.startswith("https://"):
        return "не подтверждён"
    if raw or (title or "").strip():
        return "указан методистом"
    return "нет источника"


def _basis_lines(course: CourseDocument) -> list[str]:
    lines: list[str] = []
    seen: set[str] = set()
    for document in required_legal_basis(course.brief.program_kind):
        seen.add(document["source"])
        mark = _source_mark(document["source"], document["title"])
        lines.append(
            " · ".join(
                part
                for part in (document["title"], document["source"], f"[{mark}]")
                if part
            )
        )
    for reference in course.brief.references:
        key = reference.source or reference.title
        if key in seen:
            continue
        seen.add(key)
        mark = _source_mark(reference.source, reference.title)
        parts = [part for part in (reference.title, reference.source, f"[{mark}]") if part]
        lines.append(" · ".join(parts))
    return lines or ["Нормативное основание не указано."]


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            table.rows[row_index].cells[col_index].text = value or "—"



def _competency_names(course: CourseDocument) -> dict[str, str]:
    return {node.id: node.name for node in course.competencies.nodes}


def build_docx(course: CourseDocument) -> bytes:
    """Программа в .docx ближе к двум шаблонам АЦР (ПК и общеразвивающая)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    from .kit_content import issued_document

    brief = course.brief
    kind = brief.program_kind
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    def center(line: str) -> None:
        paragraph = doc.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def bullets(lines: list[str]) -> None:
        for line in lines:
            doc.add_paragraph(line, style="List Bullet")

    titles = {module.id: module.title for module in course.modules}
    contact = sum(
        module.lecture_hours + module.practice_hours + module.lab_hours
        for module in course.modules
    )
    self_study = sum(module.self_study_hours for module in course.modules)
    module_total = sum(module.hours for module in course.modules)

    # --- Титул (оба шаблона АЦР: ПК и общеразвивающая) ---
    if kind == "general":
        org_lines = (
            "Автономная некоммерческая организация",
            "дополнительного образования",
            "«Академия Информационных Технологий»",
        )
    else:
        org_lines = (
            "Автономная некоммерческая организация",
            "дополнительного профессионального образования",
            "«Академия Информационных Технологий»",
        )
    for line in (
        *org_lines,
        "ОГРН: ____________________  ИНН/КПП: ____________________",
        "Российская Федерация, Республика Ингушетия, г. Магас",
        "",
        "УТВЕРЖДЕНО",
        "Приказом № ____-ОБР от «___» __________ 20__ г.",
        "",
        _program_title(kind),
        f"«{brief.title or 'Без названия'}»",
        "",
        "Магас",
        "20__",
    ):
        center(line)

    center("")
    center("Реквизиты приказа организации: ____________________")
    center("Преподаватель: ____________________")
    center("Подпись: ____________________")

    # --- 1. Общая характеристика ---
    if kind == "general":
        doc.add_heading("Пояснительная записка", level=1)
        doc.add_heading("1.1. Актуальность программы", level=2)
        bullets(brief.goals or ["Цель не заполнена."])
        doc.add_heading("1.2. Категории обучающихся", level=2)
        doc.add_paragraph(brief.audience or "____________________")
        if brief.level:
            doc.add_paragraph(f"Уровень освоения: {brief.level}.")
        doc.add_heading("1.3. Цель и задачи программы", level=2)
        bullets(brief.goals or ["Цель не заполнена."])
        doc.add_heading("1.4. Сроки реализации и объём", level=2)
    else:
        doc.add_heading("1. Общая характеристика программы", level=1)
        doc.add_heading("1.1. Цель реализации программы", level=2)
        bullets(brief.goals or ["Цель не заполнена."])
        doc.add_heading("1.2. Требования к слушателям", level=2)
        doc.add_paragraph(brief.audience or "____________________")
        if brief.level:
            doc.add_paragraph(f"Входной уровень: {brief.level}.")
        doc.add_heading("1.3. Форма обучения", level=2)
        doc.add_paragraph(FORMAT_LABELS.get(brief.format, brief.format or "____________________"))
        doc.add_heading("1.4. Объём программы", level=2)

    bullets(
        [
            f"Трудоёмкость: {hours(brief.duration_hours)} академических часов.",
            f"Контактная работа: {hours(contact)} ч.",
            f"Самостоятельная работа слушателей (СРС): {hours(self_study)} ч.",
            (
                f"Итоговая аттестация: "
                f"{brief.attestation_form or '____________________'}, "
                f"{hours(brief.attestation_hours)} ч."
            ),
            f"Сумма модулей: {hours(module_total)} ч. "
            f"(вместе с аттестацией = {hours(module_total + brief.attestation_hours)} ч.).",
        ]
    )

    doc.add_heading(
        "1.5. Выдаваемый документ" if kind != "general" else "1.5. Документ по итогам обучения",
        level=2,
    )
    issued = issued_document(course) or (
        "Выдаваемый документ: ____________________"
    )
    doc.add_paragraph(issued)

    doc.add_heading("1.6. Нормативные основания", level=2)
    bullets(_basis_lines(course))

    # --- Карта компетенций / ЗУН ---
    if kind == "qualification":
        heading_comp = "1.7. Планируемые результаты обучения (компетенции)"
    else:
        heading_comp = "1.7. Планируемые результаты обучения"
    doc.add_heading(heading_comp, level=2)
    nodes = course.competencies.nodes
    by_name = _competency_names(course)
    if kind == "qualification":
        _add_table(
            doc,
            [
                "Компетенция",
                "Трудовая функция",
                "Знания",
                "Умения",
                "Навыки",
                "Модули",
            ],
            [
                [
                    node.name,
                    node.labor_function_code or "—",
                    "; ".join(node.knowledge) or "—",
                    "; ".join(node.skills) or "—",
                    "; ".join(node.abilities) or "—",
                    "; ".join(titles.get(mid, mid) for mid in node.module_ids) or "—",
                ]
                for node in nodes
            ]
            or [["—", "—", "—", "—", "—", "—"]],
        )
    else:
        _add_table(
            doc,
            ["Результат / компетенция", "Знания", "Умения", "Навыки", "Модули"],
            [
                [
                    node.name,
                    "; ".join(node.knowledge) or "—",
                    "; ".join(node.skills) or "—",
                    "; ".join(node.abilities) or "—",
                    "; ".join(titles.get(mid, mid) for mid in node.module_ids) or "—",
                ]
                for node in nodes
            ]
            or [["—", "—", "—", "—", "—"]],
        )

    # --- 2. Учебный план ---
    doc.add_heading("2. Учебный план", level=1)
    plan_rows = [
        [
            str(index + 1),
            module.title,
            hours(module.hours),
            hours(module.lecture_hours),
            hours(module.practice_hours),
            hours(module.lab_hours),
            hours(module.self_study_hours),
        ]
        for index, module in enumerate(course.modules)
    ]
    plan_rows.append(
        [
            "",
            "Итоговая аттестация",
            hours(brief.attestation_hours),
            "0",
            "0",
            "0",
            "0",
        ]
    )
    plan_rows.append(
        [
            "",
            "Итого",
            hours(module_total + brief.attestation_hours),
            hours(sum(m.lecture_hours for m in course.modules)),
            hours(sum(m.practice_hours for m in course.modules)),
            hours(sum(m.lab_hours for m in course.modules)),
            hours(self_study),
        ]
    )
    _add_table(doc, ["№", "Наименование", "Всего", "ЛК", "ПР", "ЛБ", "СРС"], plan_rows)

    # --- 2.1 Учебно-тематический план ---
    doc.add_heading("2.1. Учебно-тематический план", level=2)
    topic_rows: list[list[str]] = []
    for module_index, module in enumerate(course.modules, start=1):
        plan = list(module.plan or [])
        if not plan and module.topics:
            # fallback: topics without hour split still appear
            for topic in module.topics:
                topic_rows.append(
                    [
                        str(module_index),
                        module.title,
                        topic,
                        hours(0),
                        hours(0),
                        hours(0),
                        hours(0),
                        hours(0),
                    ]
                )
            continue
        if not plan:
            topic_rows.append(
                [
                    str(module_index),
                    module.title,
                    module.title,
                    hours(module.hours),
                    hours(module.lecture_hours),
                    hours(module.practice_hours),
                    hours(module.lab_hours),
                    hours(module.self_study_hours),
                ]
            )
            continue
        for topic in plan:
            topic_hours = (
                topic.lecture_hours
                + topic.practice_hours
                + topic.lab_hours
                + topic.self_study_hours
            )
            topic_rows.append(
                [
                    str(module_index),
                    module.title,
                    topic.title,
                    hours(topic_hours),
                    hours(topic.lecture_hours),
                    hours(topic.practice_hours),
                    hours(topic.lab_hours),
                    hours(topic.self_study_hours),
                ]
            )
    _add_table(
        doc,
        ["№ мод.", "Модуль", "Тема", "Всего", "ЛК", "ПР", "ЛБ", "СРС"],
        topic_rows or [["—", "—", "—", "—", "—", "—", "—", "—"]],
    )

    doc.add_heading("2.2. Результаты обучения по модулям", level=2)
    for module in course.modules:
        doc.add_heading(module.title, level=3)
        for outcome in module.learning_outcomes:
            doc.add_paragraph(outcome, style="List Bullet")
        linked = [
            node.name
            for node in nodes
            if module.id in (node.module_ids or [])
            or module.id in (getattr(module, "competency_ids", None) or [])
        ]
        # also via module.competency_ids
        by_id = {node.id: node.name for node in nodes}
        for cid in module.competency_ids or []:
            name = by_id.get(cid)
            if name and name not in linked:
                linked.append(name)
        if linked:
            doc.add_paragraph("Связанные компетенции: " + "; ".join(linked))

    doc.add_heading("Паспорт результата", level=2)
    doc.add_paragraph(
        "Часы, вид программы и состав модулей считает код. "
        "Формулировки лекций в этот расчёт не входят."
    )
    doc.add_heading("Почему план такой", level=2)
    doc.add_paragraph(
        "Формулировки лекций при новом прогоне модели могут отличаться. "
        "Часы, результаты и связь со средствами проверки считает код."
    )
    for row in module_rationales(course.as_api()):
        doc.add_paragraph(row["text"], style="List Bullet")

    # --- 3. Итоговая аттестация ↔ компетенции ---
    doc.add_heading("3. Итоговая аттестация", level=1)
    doc.add_paragraph(
        f"Форма: {brief.attestation_form or '____________________'}; "
        f"объём: {hours(brief.attestation_hours)} ч."
    )
    doc.add_paragraph(
        "Связь компетенций со средствами проверки (тест / задание / итоговая аттестация):"
    )
    by_name = _competency_names(course)
    assess_by_comp: dict[str, int] = {}
    task_by_comp: dict[str, int] = {}
    for assessment in course.assessments:
        for question in assessment.questions:
            cid = question.competency_id or ""
            if cid:
                assess_by_comp[cid] = assess_by_comp.get(cid, 0) + 1
    for seminar in course.seminars:
        for task in seminar.tasks:
            cid = task.competency_id or ""
            if cid:
                task_by_comp[cid] = task_by_comp.get(cid, 0) + 1
    matrix_rows: list[list[str]] = []
    for node in nodes:
        qn = assess_by_comp.get(node.id, 0)
        tn = task_by_comp.get(node.id, 0)
        # fallback: module-level link counts as covered if questions/tasks exist on those modules
        if qn == 0:
            mod_ids = set(node.module_ids or [])
            for assessment in course.assessments:
                if assessment.module_id in mod_ids:
                    qn += len(assessment.questions)
        if tn == 0:
            mod_ids = set(node.module_ids or [])
            for seminar in course.seminars:
                if seminar.module_id in mod_ids:
                    tn += len(seminar.tasks)
        covered = "да" if (qn and tn) else ("частично" if (qn or tn) else "нет")
        matrix_rows.append(
            [
                node.name,
                str(qn),
                str(tn),
                brief.attestation_form or "—",
                covered,
            ]
        )
    _add_table(
        doc,
        ["Компетенция / результат", "Вопросов теста", "Заданий", "Итоговая аттестация", "Покрыто"],
        matrix_rows or [["—", "0", "0", "—", "нет"]],
    )
    hours_match = abs(module_total + brief.attestation_hours - brief.duration_hours) < 0.001
    doc.add_paragraph(
        "Автопроверка часов: "
        + (
            f"сходятся ({hours(module_total + brief.attestation_hours)} = "
            f"{hours(brief.duration_hours)})."
            if hours_match
            else (
                f"расхождение — модули {hours(module_total)} + аттестация "
                f"{hours(brief.attestation_hours)} ≠ {hours(brief.duration_hours)}."
            )
        )
    )

    # --- Приложения ---
    doc.add_heading("Приложение 1. Оценочные материалы", level=1)
    questions = [
        (assessment.module_id, question)
        for assessment in course.assessments
        for question in assessment.questions
    ]
    if not questions:
        doc.add_paragraph("Вопросы заполняются после утверждения каркаса.")
    for index, (module_id, question) in enumerate(questions, start=1):
        doc.add_paragraph(f"{index}. {question.stem}")
        doc.add_paragraph(f"Модуль: {titles.get(module_id, module_id)}")
        if question.competency_id:
            qname = by_name.get(question.competency_id, question.competency_id)
            doc.add_paragraph(f"Компетенция: {qname}")
        for option in question.options or []:
            doc.add_paragraph(option, style="List Bullet")
        doc.add_paragraph(f"Ключ: {question.answer}")
        doc.add_paragraph(f"Результат обучения: {question.learning_outcome}")

    doc.add_heading("Приложение 2. Практические задания", level=1)
    tasks = [
        (seminar.module_id, task)
        for seminar in course.seminars
        for task in seminar.tasks
    ]
    if not tasks:
        doc.add_paragraph("Задания заполняются после утверждения каркаса.")
    for module_id, task in tasks:
        doc.add_heading(f"{task.title} · {hours(task.time_min)} мин", level=2)
        doc.add_paragraph(f"Модуль: {titles.get(module_id, module_id)}")
        if task.competency_id:
            tname = by_name.get(task.competency_id, task.competency_id)
            doc.add_paragraph(f"Компетенция: {tname}")
        doc.add_paragraph(task.prompt)
        for step in task.steps:
            doc.add_paragraph(step, style="List Number")
        doc.add_paragraph("Критерии")
        for criterion in task.criteria:
            doc.add_paragraph(criterion, style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()



def _rgb(value: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(value)


def _paint(slide: Any) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb("F6F4F1")


def _rule(slide: Any, left: float, top: float, width: float = 2.2) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Emu(32000)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb("6E2432")
    shape.line.fill.background()


def _add_textbox(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: int = 20,
    bold: bool = False,
    color: str = "1C1614",
    font: str = "Calibri",
) -> None:
    from pptx.util import Inches, Pt as PptPt

    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = font


def _add_bullets(slide: Any, items: list[str]) -> None:
    from pptx.util import Inches, Pt as PptPt

    box = slide.shapes.add_textbox(Inches(0.85), Inches(2.35), Inches(11.5), Inches(4.5))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items[:6]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"•  {item}"
        paragraph.level = 0
        paragraph.space_after = PptPt(12)
        paragraph.font.size = PptPt(20)
        paragraph.font.color.rgb = _rgb("1C1614")
        paragraph.font.name = "Calibri"


def _new_presentation():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    return presentation


def _presentation_bytes(presentation) -> bytes:
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _add_course_cover(presentation, course: CourseDocument) -> None:
    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    _paint(cover)
    _add_textbox(cover, 0.75, 1.35, 11.5, 0.35, "СТУДИЯ МЕТОДИСТА", 14, False, "6E2432", "Calibri")
    _add_textbox(
        cover,
        0.75,
        1.85,
        11.5,
        1.6,
        course.brief.title or "Программа",
        40,
        False,
        "1C1614",
        "Georgia",
    )
    _rule(cover, 0.75, 3.6, 2.4)
    audience = course.brief.audience.strip()
    duration = course.brief.duration_hours
    subtitle = " · ".join(
        part
        for part in (
            audience,
            f"{duration:g} ч" if duration else "",
            "Материалы лекций",
        )
        if part
    )
    _add_textbox(cover, 0.75, 3.85, 11.5, 0.6, subtitle, 18, False, "5C534C")


def _add_lecture_slides(presentation, lecture, module_title: str) -> None:
    blank = presentation.slide_layouts[6]
    opener = presentation.slides.add_slide(blank)
    _paint(opener)
    _add_textbox(opener, 0.75, 2.35, 11.5, 0.35, "ЛЕКЦИЯ", 14, False, "6E2432")
    _add_textbox(opener, 0.75, 2.85, 11.5, 1.4, module_title, 36, False, "1C1614", "Georgia")
    _rule(opener, 0.75, 4.45, 2.2)
    if not lecture.slides:
        empty = presentation.slides.add_slide(blank)
        _paint(empty)
        _add_textbox(empty, 0.8, 3.1, 11.7, 1, "Слайды этой лекции ещё не заполнены.", 22, False, "5C534C", "Georgia")
        return
    for index, slide in enumerate(lecture.slides, start=1):
        page = presentation.slides.add_slide(blank)
        _paint(page)
        _add_textbox(page, 0.75, 0.38, 10.4, 0.32, module_title, 13, False, "6E2432")
        _add_textbox(page, 11.2, 0.38, 1.3, 0.32, f"{index} / {len(lecture.slides)}", 13, False, "5C534C")
        _add_textbox(page, 0.75, 0.9, 11.7, 0.9, slide.title, 32, False, "1C1614", "Georgia")
        _rule(page, 0.75, 1.9, 1.8)
        _add_bullets(page, list(slide.bullets))


def build_lecture_pptx(course: CourseDocument, lecture) -> bytes:
    presentation = _new_presentation()
    module_title = next(
        (module.title for module in course.modules if module.id == lecture.module_id),
        lecture.module_id,
    )
    _add_lecture_slides(presentation, lecture, module_title)
    return _presentation_bytes(presentation)


def build_lecture_docx(course: CourseDocument, lecture) -> bytes:
    from docx import Document
    from docx.shared import Pt

    module_title = next(
        (module.title for module in course.modules if module.id == lecture.module_id),
        lecture.module_id,
    )
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    doc.add_heading(module_title, level=1)
    doc.add_paragraph(f"Курс: {course.brief.title or 'Без названия'}")
    doc.add_paragraph(f"Длительность: {hours(lecture.duration_min)} мин.")
    doc.add_heading("План занятия", level=2)
    for item in lecture.outline or ["План не заполнен."]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Ход лекции", level=2)
    if not lecture.slides:
        doc.add_paragraph("Слайды не заполнены.")
    for index, slide in enumerate(lecture.slides, start=1):
        doc.add_heading(f"{index}. {slide.title}", level=3)
        for bullet in slide.bullets or ["—"]:
            doc.add_paragraph(bullet, style="List Bullet")
        if slide.speaker_notes:
            doc.add_paragraph(f"Заметки преподавателя: {slide.speaker_notes}")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _add_results_slide(presentation, course: CourseDocument) -> None:
    """Слайд результатов обучения / компетенций — связка с аттестацией."""
    nodes = list(course.competencies.nodes or [])
    if not nodes:
        return
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _paint(slide)
    kind = course.brief.program_kind
    heading = (
        "Планируемые результаты (компетенции)"
        if kind == "qualification"
        else "Планируемые результаты обучения"
    )
    _add_textbox(slide, 0.7, 0.45, 12, 0.7, heading, 26, True, "2C2825", "Georgia")
    bullets: list[str] = []
    for node in nodes[:8]:
        zun = []
        if node.knowledge:
            zun.append("З: " + node.knowledge[0][:60])
        if node.skills:
            zun.append("У: " + node.skills[0][:60])
        if node.abilities:
            zun.append("Н: " + node.abilities[0][:60])
        line = node.name
        if zun:
            line = f"{node.name} — {'; '.join(zun)}"
        bullets.append(line[:140])
    _add_bullets(slide, bullets)
    assessed = sum(len(a.questions) for a in course.assessments)
    tasks = sum(len(s.tasks) for s in course.seminars)
    _add_textbox(
        slide,
        0.7,
        6.4,
        12,
        0.5,
        f"Проверка: {assessed} вопросов теста, {tasks} заданий семинаров.",
        14,
        False,
        "5C534C",
        "Georgia",
    )


def build_pptx(course: CourseDocument) -> bytes:
    presentation = _new_presentation()
    _add_course_cover(presentation, course)
    _add_results_slide(presentation, course)
    if not course.lectures:
        empty = presentation.slides.add_slide(presentation.slide_layouts[6])
        _paint(empty)
        _add_textbox(empty, 0.8, 3.1, 11.7, 1, "Слайды появятся после утверждения каркаса.", 22, False, "5C534C", "Georgia")
    titles = {module.id: module.title for module in course.modules}
    for lecture in course.lectures:
        _add_lecture_slides(presentation, lecture, titles.get(lecture.module_id, lecture.module_id))
    if len(presentation.slides) < 1:
        empty = presentation.slides.add_slide(presentation.slide_layouts[6])
        _paint(empty)
    return _presentation_bytes(presentation)
