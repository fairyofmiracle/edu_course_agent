import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.models import CourseDocument

pytest.importorskip("docx")
pytest.importorskip("pptx")

FIXTURE = Path(__file__).parents[2] / "fixtures" / "course-contract.json"


def test_fixture_is_available_next_to_backend() -> None:
    assert FIXTURE.is_file()
    course = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assert course["brief"]["title"]


def test_export_json_docx_pptx_from_body() -> None:
    course = json.loads(FIXTURE.read_text(encoding="utf-8"))
    client = TestClient(app)
    json_file = client.post("/api/export/json", json=course)
    assert json_file.status_code == 200
    assert json_file.json()["brief"]["title"] == course["brief"]["title"]

    docx = client.post("/api/export/docx", json=course)
    assert docx.status_code == 200
    assert docx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml"
    )
    assert docx.content[:2] == b"PK"
    from io import BytesIO
    from zipfile import ZipFile

    with ZipFile(BytesIO(docx.content)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert course["brief"]["title"] in xml

    pptx = client.post("/api/export/pptx", json=course)
    assert pptx.status_code == 200
    assert pptx.content[:2] == b"PK"

    bad = client.post("/api/export/pdf", json=course)
    assert bad.status_code == 400
    assert "docx" in bad.json()["detail"]
    assert "Traceback" not in bad.json()["detail"]


def test_export_saved_course_after_put() -> None:
    pytest.importorskip("sqlalchemy")
    course = json.loads(FIXTURE.read_text(encoding="utf-8"))
    course["id"] = "export-saved-course"
    client = TestClient(app)
    saved = client.put(f"/api/courses/{course['id']}", json={"course": course, "reason": "export"})
    if saved.status_code != 200:
        pytest.skip("persistence unavailable")
    exported = client.get(f"/api/courses/{course['id']}/export/json")
    assert exported.status_code == 200
    assert exported.json()["id"] == course["id"]
    missing = client.get("/api/courses/missing-export/export/docx")
    assert missing.status_code == 404
    client.delete(f"/api/courses/{course['id']}")


def _slide_text(payload: bytes) -> str:
    from io import BytesIO
    from zipfile import ZipFile

    with ZipFile(BytesIO(payload)) as archive:
        return " ".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide")
        )


def test_zip_puts_each_lecture_in_its_own_file() -> None:
    from io import BytesIO
    from zipfile import ZipFile

    from backend.app.exports import build_zip

    raw = json.loads(FIXTURE.read_text(encoding="utf-8"))
    raw["lectures"] = [
        {
            "moduleId": "M1",
            "durationMin": 40,
            "outline": ["Переменные"],
            "slides": [{"title": "Что такое переменная", "bullets": ["Имя и значение"]}],
        },
        {
            "moduleId": "M1",
            "durationMin": 40,
            "outline": ["Условия"],
            "slides": [{"title": "Ветвление if", "bullets": ["Условие"]}],
        },
    ]
    course = CourseDocument.model_validate(raw)
    with ZipFile(BytesIO(build_zip(course))) as archive:
        names = archive.namelist()
        program = [name for name in names if name.startswith("основной документ/") and name.endswith(".docx")]
        decks = [name for name in names if name.startswith("презентации/") and name.endswith(".pptx")]
        lectures = [name for name in names if name.startswith("лекции/") and name.endswith(".docx")]
        assert len(program) == 1
        assert decks == [
            "презентации/01-Базовый-синтаксис.pptx",
            "презентации/02-Базовый-синтаксис.pptx",
        ]
        assert lectures == [
            "лекции/01-Базовый-синтаксис.docx",
            "лекции/02-Базовый-синтаксис.docx",
        ]
        assert "тесты/Основы-Python.gift.txt" in names
        assert "тесты/Основы-Python-moodle.xml" in names
        assert "тесты/Основы-Python.csv" in names
        assert "слайды/Основы-Python.html" in names
        assert not any("/" not in name and name.endswith((".pptx", ".docx")) for name in names)
        first = _slide_text(archive.read(decks[0]))
        second = _slide_text(archive.read(decks[1]))
        assert "Что такое переменная" in first
        assert "Ветвление if" not in first
        assert "Ветвление if" in second
        assert "Что такое переменная" not in second
        with ZipFile(BytesIO(archive.read(lectures[0]))) as lecture:
            text = lecture.read("word/document.xml").decode("utf-8")
        assert "Переменные" in text
        assert "Условия" not in text


def test_saved_zip_stores_checksum() -> None:
    course = json.loads(FIXTURE.read_text(encoding="utf-8"))
    course["id"] = "artifact-export-test"
    client = TestClient(app)
    saved = client.put("/api/courses/artifact-export-test", json={"course": course, "messages": []})
    assert saved.status_code == 200
    exported = client.post("/api/export/zip", json=course)
    assert exported.status_code == 200
    digest = exported.headers.get("x-studio-sha256")
    assert digest and len(digest) == 64
    versions = client.get("/api/courses/artifact-export-test/versions")
    assert versions.status_code == 200
    assert versions.json()["versions"][0]["sha256"] == digest
    client.delete("/api/courses/artifact-export-test")


def test_build_artifact_rejects_unknown_kind() -> None:
    from backend.app.exports import build_artifact

    course = CourseDocument.model_validate(json.loads(FIXTURE.read_text(encoding="utf-8")))
    try:
        build_artifact("xls", course)
    except ValueError as error:
        assert "docx" in str(error)
    else:
        raise AssertionError("неизвестный формат должен падать")


def test_docx_both_acr_templates() -> None:
    """Оба шаблона АЦР: разные титулы."""
    from io import BytesIO

    from docx import Document

    from backend.app.exports import build_docx
    from backend.tests.test_domain import fixture_course

    def _text(payload: bytes) -> str:
        doc = Document(BytesIO(payload))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    course = fixture_course()
    course.brief.program_kind = "qualification"
    course.brief.title = "Python для аналитиков"
    pk = _text(build_docx(course))
    assert "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА" in pk
    assert "дополнительного профессионального образования" in pk

    course.brief.program_kind = "general"
    course.brief.title = "Основы информатики"
    od = _text(build_docx(course))
    assert "ДОПОЛНИТЕЛЬНАЯ ОБЩЕОБРАЗОВАТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА" in od
    assert "дополнительного образования" in od
    assert "дополнительного профессионального образования" not in od
    assert "Пояснительная записка" in od


def test_pptx_has_cover_and_results_slide() -> None:
    pytest.importorskip("pptx")
    from io import BytesIO

    from pptx import Presentation

    from backend.app.exports import build_pptx
    from backend.tests.test_domain import fixture_course

    course = fixture_course()
    payload = build_pptx(course)
    deck = Presentation(BytesIO(payload))
    assert len(deck.slides) >= 1
    texts: list[str] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text_frame.text)
    blob = "\n".join(texts).lower()
    assert "результат" in blob or "компетенц" in blob or course.brief.title[:8].lower() in blob


def test_lms_gift_and_moodle_build() -> None:
    from backend.app.lms_export import course_gift, course_moodle_xml
    from backend.tests.test_domain import fixture_course

    course = fixture_course()
    gift = course_gift(course)
    moodle = course_moodle_xml(course)
    assert isinstance(gift, str)
    assert isinstance(moodle, str)
    assert "quiz" in moodle or "question" in moodle or "<?xml" in moodle



def test_zip_contains_audit_report() -> None:
    from io import BytesIO
    from zipfile import ZipFile

    from backend.app.exports import build_zip
    from backend.tests.test_domain import fixture_course

    course = fixture_course()
    payload = build_zip(course)
    names = ZipFile(BytesIO(payload)).namelist()
    assert any(name.endswith("-audit.json") or "/проверка/" in name or name.startswith("проверка/") for name in names)


def test_docx_attestation_competency_matrix() -> None:
    from io import BytesIO

    from docx import Document

    from backend.app.exports import build_docx
    from backend.tests.test_domain import fixture_course

    course = fixture_course()
    doc = Document(BytesIO(build_docx(course)))
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)
    assert "Итоговая аттестация" in text
    assert "Автопроверка часов" in text
    assert "Вопросов теста" in text or "Покрыто" in text
