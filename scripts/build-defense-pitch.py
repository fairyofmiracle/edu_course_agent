"""Питч защиты 30 октября. Запуск: python scripts/build-defense-pitch.py"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).parents[1] / "docs" / "pitch-zashchita.pptx"

INK = RGBColor(0x14, 0x13, 0x11)
PAPER = RGBColor(0xF7, 0xF6, 0xF2)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x0C, 0x7A, 0x4B)
MUTED = RGBColor(0x5E, 0x5E, 0x66)
LINE = RGBColor(0xE6, 0xE6, 0xEA)


def run(paragraph, text: str, size: int, bold: bool, color: RGBColor) -> None:
    piece = paragraph.add_run()
    piece.text = text
    piece.font.size = Pt(size)
    piece.font.bold = bold
    piece.font.color.rgb = color
    piece.font.name = "Calibri"


def paint(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def box(slide, left, top, width, height, text, size=18, bold=False, color=INK):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run(paragraph, text, size, bold, color)
    return shape


def bullets(slide, left, top, width, height, lines: list[str], size=20) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.level = 0
        paragraph.space_after = Pt(10)
        run(paragraph, line, size, False, INK)


def card(slide, left, top, width, height) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.08


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def kicker(slide, text: str) -> None:
    box(slide, 0.6, 0.28, 12, 0.3, text, 13, True, GREEN)


def title(slide, text: str) -> None:
    box(slide, 0.6, 0.58, 12, 0.9, text, 32, True)


def build() -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    slide = deck.slides.add_slide(blank)
    paint(slide)
    box(slide, 0.7, 2.15, 11, 0.35, "СТУДИЯ МЕТОДИСТА", 14, True, GREEN)
    box(slide, 0.7, 2.55, 11.5, 1.4, "Методий", 60, True)
    box(slide, 0.7, 4.3, 10, 1.2, "Методист приносит сырой запрос.\nНа выходе документ, который можно сдать в нормоконтроль.", 22, False, MUTED)
    notes(slide, "Не чат, который печатает программу. Место, где методист собирает программу и может её подписать.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "01  Где теряется время")
    title(slide, "Текст есть. Подписать его нельзя.")
    lines = [
        "«Месяц» в запросе превращается в выдуманные часы.",
        "Приказ и код профстандарта появляются, хотя их не было в материалах.",
        "Весь курс пишется сразу, и методист сверяет простыню вручную.",
    ]
    for index, line in enumerate(lines):
        top = 2.0 + index * 1.35
        card(slide, 0.6, top, 12.1, 1.2)
        box(slide, 0.9, top + 0.32, 11.4, 0.6, line, 22, True)
    notes(slide, "Боль не в том, что текста мало. Боль в том, что текст нельзя подписать. Здесь место для QR на студию.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "02  Идея")
    title(slide, "Сначала каркас. Комплект — после утверждения.")
    bullets(slide, 0.7, 1.9, 11.5, 4.5, [
        "Слева программа, справа диалог.",
        "Модель предлагает формулировки. Часы, шаблон и нормы держит код.",
        "«Утвердить каркас» появляется, только когда обязательные данные на месте.",
        "Снятие утверждения возвращает правку каркаса.",
        "Лекции, семинары и тесты не пишутся, пока каркас не утверждён.",
    ], 24)
    notes(slide, "Методист остаётся тем, кто решает. Токены на комплект не тратятся до его кнопки.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "03  Как устроено")
    title(slide, "Один ход — один инструмент.")
    blocks = [
        ("Экран", "Next.js"),
        ("Оркестрация", "FastAPI, один клиент к шлюзу"),
        ("Шлюзы", "Только Cloud.ru и GigaSchool"),
        ("Память", "Версии курса, конфликт не затирает правку"),
        ("Выход", "DOCX, PPTX и JSON, не скриншот чата"),
    ]
    for index, (head, body) in enumerate(blocks):
        left = 0.55 + (index % 3) * 4.15
        top = 1.9 if index < 3 else 4.35
        card(slide, left, top, 3.95, 2.1)
        box(slide, left + 0.25, top + 0.35, 3.45, 0.4, head, 14, True, GREEN)
        box(slide, left + 0.25, top + 0.9, 3.45, 0.9, body, 20, True)
    notes(slide, "В модель уходит снимок шага, а не весь курс. Чужой адрес шлюза обрывается до запроса.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "04  Проверка")
    title(slide, "Вердикт до выгрузки")
    checks = [
        "Часы только из слов методиста. «Месяц» — срок, не объём.",
        "Сумма часов модулей сходится с объёмом программы.",
        "Общеразвивающая — приказ № 629. Повышение квалификации — 273-ФЗ и приказ № 266.",
        "Номер приказа и код трудовой функции не выдумываются.",
        "У вопроса есть ключ и результат модуля. Практика закрывает тот же результат.",
    ]
    for index, line in enumerate(checks):
        top = 1.75 + index * 0.95
        box(slide, 0.7, top, 0.45, 0.4, "●", 18, True, GREEN)
        box(slide, 1.2, top, 11, 0.7, line, 20)
    notes(slide, "Жюри смотрит, можно ли документу доверять. Вкладка «Проверка» — до выгрузки.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "05  Демонстрация")
    title(slide, "Один сценарий на данных кейса")
    steps = [
        ("1", "Пустой курс. Лекции и тесты закрыты."),
        ("2", "«Примерно на месяц». Число часов само не появляется."),
        ("3", "Бриф и каркас. Утверждение — когда данные на месте."),
        ("4", "Комплект: тест, практика, слайды урока."),
        ("5", "Проверка и выгрузка DOCX / PPTX."),
    ]
    for index, (num, line) in enumerate(steps):
        top = 1.75 + index * 0.95
        card(slide, 0.6, top, 12.1, 0.82)
        box(slide, 0.85, top + 0.18, 0.6, 0.45, num, 22, True, GREEN)
        box(slide, 1.6, top + 0.2, 10.6, 0.45, line, 22)
    notes(slide, "В зале показываем вживую. Если сеть упадёт — запись экрана. Не читать слайд.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "06  Польза")
    title(slide, "Что методист перестаёт делать руками")
    bullets(slide, 0.7, 1.9, 11.5, 4.2, [
        "Не вылавливает выдуманный объём.",
        "Не ищет в тексте приказ, которого не было в источниках.",
        "Не платит токенами за лекции до согласования каркаса.",
        "Получает файл, который открывается в Word и PowerPoint.",
    ], 26)
    box(slide, 0.7, 6.3, 11.5, 0.6, "Процент экономии назовём после замера. Логика уже в продукте.", 16, False, MUTED)
    notes(slide, "Не говорить выдуманный процент. Говорить: нет утверждённого каркаса — нет комплекта.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "07  Внедрение")
    title(slide, "Что нужно, чтобы это поставить")
    bullets(slide, 0.7, 1.85, 11.5, 4.8, [
        "Данные: бриф, материалы методиста, ссылки на официальное опубликование.",
        "Технология: сервер студии и разрешённый шлюз.",
        "Если шлюз молчит, текст модулей не выдумывается. Явные часы всё равно записываются.",
        "Применимость нормы подтверждает методист, не страница приказа.",
        "Дальше: замер токенов и пилот на одном курсе.",
    ], 22)
    notes(slide, "Не обещаем заменить методиста. Обещаем не подписать за него чужие часы и чужой приказ.")

    slide = deck.slides.add_slide(blank)
    paint(slide)
    box(slide, 0.8, 2.5, 11.5, 1.8, "Модель пишет.\nКод не выпускает.", 48, True)
    box(slide, 0.8, 5.2, 11, 0.8, "Проверка до выгрузки. Готовы показать ещё раз на вопросах.", 20, False, MUTED)
    notes(slide, "Закрытие. Одна фраза и пауза.")

    deck.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
