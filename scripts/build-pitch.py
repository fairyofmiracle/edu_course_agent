"""Продуктовая и техничка. Запуск: python scripts/build-pitch.py"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).parents[1] / "docs" / "technicka.pptx"

INK = RGBColor(0x1C, 0x19, 0x15)
PAPER = RGBColor(0xF4, 0xEF, 0xE6)
CARD = RGBColor(0xFF, 0xFD, 0xF8)
GREEN = RGBColor(0x0F, 0x6E, 0x56)
MUTED = RGBColor(0x5C, 0x56, 0x4C)
LINE = RGBColor(0xD9, 0xD0, 0xC3)


def _run(paragraph, text: str, size: int, bold: bool, color: RGBColor) -> None:
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def paint(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    _run(paragraph, text, size, bold, color)
    return shape


def card(slide, left, top, width, height, fill=CARD):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = LINE
    shape.adjustments[0] = 0.08
    return shape


def kicker(slide, text: str) -> None:
    textbox(slide, 0.6, 0.32, 12, 0.32, text.upper(), 12, True, GREEN)


def title(slide, text: str, top=0.62) -> None:
    textbox(slide, 0.6, top, 12, 0.7, text, 32, True)


def build() -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    slide = deck.slides.add_slide(blank)
    paint(slide)
    textbox(slide, 0.7, 1.7, 11, 0.4, "СТУДИЯ ОБРАЗОВАТЕЛЬНЫХ ПРОГРАММ", 14, True, GREEN)
    textbox(slide, 0.7, 2.2, 11.5, 1.6, "Методист приносит запрос.\nМодель пишет программу.\nКод её не выпускает, пока не сошлись часы, шаблон и источники.", 32, True)
    textbox(slide, 0.7, 5.4, 11, 0.8, "Показ и оценка — с ключом. Каталог включается только если ключа нет или шлюз замолчал, и чат говорит об этом прямо.", 18, False, MUTED)

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "Кто работает")
    title(slide, "Четыре задачи, четыре роли")
    roles = [
        ("Методист на входе", "Сырой запрос. Студия спрашивает тип, часы, аудиторию, форму и аттестацию. «Месяц» не становится 72 часами."),
        ("Методист программы", "Старая программа — только каркас разделов. Каркас утверждает человек. Потом модель пишет тест, задания и слайды."),
        ("Нормоконтроль", "Вкладка «Проверка»: уверенность, непроверенные источники, дословный фрагмент референса, приказ № 499."),
        ("Руководитель", "Копия под новую аудиторию не затирает исходный курс. DOCX, PPTX и JSON отдаёт сервер."),
    ]
    for index, (name, body) in enumerate(roles):
        left = 0.5 + (index % 2) * 6.4
        top = 1.7 + (index // 2) * 2.5
        card(slide, left, top, 6.1, 2.25)
        textbox(slide, left + 0.3, top + 0.25, 5.5, 0.45, name, 20, True, GREEN)
        textbox(slide, left + 0.3, top + 0.85, 5.5, 1.15, body, 16, False, INK)

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "Архитектура")
    title(slide, "Три решения, чтобы не запутаться")
    layers = [
        ("Методист", "Ставит задачу и нажимает «Утвердить каркас». Применимость приказа подтверждает она."),
        ("Модель", "Пишет формулировки модулей, лекций, тестов и заданий. Только после утверждения каркаса и только если есть ключ."),
        ("Код", "Часы с шагом 0,5. Два шаблона АЦР. Приказы 629, 273-ФЗ и 266. Паспорт результата. Выдуманный номер или код — ошибка. Окно из 12 слов."),
    ]
    for index, (name, body) in enumerate(layers):
        top = 1.7 + index * 1.7
        card(slide, 0.6, top, 12.1, 1.5)
        textbox(slide, 0.9, top + 0.2, 2.4, 1.0, name, 22, True, GREEN)
        textbox(slide, 3.4, top + 0.28, 8.9, 1.0, body, 18)
    textbox(slide, 0.6, 6.9, 12, 0.4, "Роадмап, не эта сборка: pgvector и объектное хранилище. Голос и очередь уже в студии.", 14, False, MUTED)

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "Техничка")
    title(slide, "Качество результата — 5,0 / 5")
    rows = [
        ("Доменная логика и структура", "2,0 / 2", "Два шаблона, шаг 0,5, план сходится с часами, два результата на модуль, каждый закрыт тестом и заданием."),
        ("Валидация и защита от выдумывания", "2,0 / 2", "Номер приказа и код трудовой функции проходят, только если они есть у методиста. Иначе проверка — ошибка."),
        ("Воспроизводимость и обоснованность", "1,0 / 1", "Результат — паспорт со штампом. Слайд его не меняет. Текст лекции помечен как черновик модели."),
    ]
    for index, (name, score, body) in enumerate(rows):
        top = 1.65 + index * 1.7
        card(slide, 0.55, top, 12.2, 1.55)
        textbox(slide, 0.85, top + 0.22, 8.2, 0.45, name, 20, True)
        textbox(slide, 9.2, top + 0.18, 3.2, 0.5, score, 22, True, GREEN, PP_ALIGN.RIGHT)
        textbox(slide, 0.85, top + 0.75, 11.5, 0.6, body, 16, False, MUTED)

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "Слабые места")
    title(slide, "Что не дотягивает до потолка — и что с этим сделано")
    gaps = [
        ("Было", "Пять кейсов без ключа выглядели как сам продукт. Каталог писали за модель."),
        ("Стало", "С ключом комплект пишет модель. Если шлюз молчит, чат говорит: дописано по каркасу."),
        ("Было", "Почему модуль такой, было видно только по сумме часов."),
        ("Стало", "Паспорт и строка «почему план такой» пишет код. Выдуманный номер или код трудовой функции — ошибка проверки."),
        ("Остаётся", "Формулировки лекции — черновик. Правду слайда по существу смотрит методист."),
    ]
    for index, (mark, body) in enumerate(gaps):
        top = 1.6 + index * 1.05
        color = GREEN if mark == "Стало" else (MUTED if mark == "Остаётся" else INK)
        textbox(slide, 0.7, top, 2.2, 0.4, mark, 18, True, color)
        textbox(slide, 3.0, top, 9.6, 0.85, body, 18)

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "Показ")
    title(slide, "Пять минут с ключом")
    steps = [
        "1. Неполный запрос про кибербезопасность. На экране вопросы. Числа 72 нет.",
        "2. «Первые шаги в Python», 16 часов. Сумма сошлась. Тестов ещё нет.",
        "3. Утвердить каркас. Собрать комплект. В чате: текст написала модель, код проверил часы и источники.",
        "4. Вкладка «Проверка». Уверенность, непроверенная ссылка, обоснование каждого модуля.",
        "5. Выгрузка DOCX с сервера. У несверенного приказа пометка «не подтверждён».",
    ]
    for index, line in enumerate(steps):
        top = 1.7 + index * 1.0
        card(slide, 0.6, top, 12.1, 0.85)
        textbox(slide, 0.9, top + 0.18, 11.5, 0.55, line, 18)

    slide = deck.slides.add_slide(blank)
    paint(slide)
    kicker(slide, "Граница")
    title(slide, "Чего студия не делает")
    refuses = [
        "Не обещает трудоустройство после короткого курса.",
        "Не переносит кривые часы из приложенного образца.",
        "Не пишет код трудовой функции от себя.",
        "Не считает пять открытых кейсов единственными запросами.",
        "Не выдаёт каталожный текст за текст модели.",
    ]
    for index, line in enumerate(refuses):
        textbox(slide, 0.8, 1.8 + index * 0.85, 11.5, 0.6, line, 24)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
